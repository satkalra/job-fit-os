import io

from docx import Document

from app.schemas.resume import EducationEntry, JobEntry, RewrittenResume
from app.services.docx_builder import build_resume_docx


def make_resume():
    return RewrittenResume(
        summary="Great engineer.",
        jobs=[
            JobEntry(
                title="Senior ML Engineer",
                company="Anthropic",
                location="San Francisco, CA",
                dates="2022 - Present",
                bullets=["Built stuff"],
            )
        ],
        skills=["Python", "SQL"],
        education=[
            EducationEntry(
                degree="B.S. Computer Science",
                school="State University",
                location="Austin, TX",
                dates="2018",
            )
        ],
    )


def test_build_returns_bytes():
    result = build_resume_docx(make_resume(), "Jane Doe", "jane@example.com")
    assert isinstance(result, bytes)
    assert len(result) > 100


def test_docx_contains_name():
    result = build_resume_docx(make_resume(), "Jane Doe", "jane@example.com")
    doc = Document(io.BytesIO(result))
    assert "Jane Doe" in " ".join(p.text for p in doc.paragraphs)


def test_docx_contains_job_title():
    result = build_resume_docx(make_resume(), "Jane Doe", "jane@example.com")
    doc = Document(io.BytesIO(result))
    text = " ".join(p.text for p in doc.paragraphs)
    assert "Senior ML Engineer" in text

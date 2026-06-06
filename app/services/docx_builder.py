import io
import logging

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from app.schemas.resume import RewrittenResume

logger = logging.getLogger(__name__)


def build_resume_docx(resume: RewrittenResume, name: str, contact: str = "") -> bytes:
    logger.debug(
        "build_resume_docx start name_present=%s contact_present=%s jobs=%s projects=%s skills=%s education=%s",
        bool(name),
        bool(contact),
        len(resume.jobs),
        len(resume.projects),
        len(resume.skills),
        len(resume.education),
    )
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    name_para = doc.add_paragraph()
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_para.add_run(name or "[unknown]")
    name_run.bold = True
    name_run.font.size = Pt(16)
    name_run.font.name = "Calibri"

    if contact:
        contact_para = doc.add_paragraph()
        contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run(contact_para.add_run(contact), size=10)

    _add_section(doc, "SUMMARY")
    _add_text(doc, resume.summary)

    _add_section(doc, "EXPERIENCE")
    for job in resume.jobs:
        _warn_missing("job", job.model_dump(), ["title", "company", "location", "dates"])
        _add_entry_heading(doc, [job.title, job.company, job.location, job.dates])
        _add_bullets(doc, job.bullets)

    if resume.projects:
        _add_section(doc, "PROJECTS")
        for project in resume.projects:
            _warn_missing("project", project.model_dump(), ["name", "dates"])
            _add_entry_heading(doc, [project.name, project.dates])
            _add_bullets(doc, project.bullets)

    _add_section(doc, "SKILLS")
    _add_text(doc, ", ".join(resume.skills))

    _add_section(doc, "EDUCATION")
    for education in resume.education:
        _warn_missing("education", education.model_dump(), ["degree", "school", "location", "dates"])
        _add_entry_heading(doc, [education.degree, education.school, education.location, education.dates])

    buf = io.BytesIO()
    doc.save(buf)
    result = buf.getvalue()
    logger.debug("build_resume_docx complete bytes=%s", len(result))
    return result


def _add_section(doc: Document, title: str) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(8)
    para.paragraph_format.space_after = Pt(2)
    run = para.add_run(title)
    run.bold = True
    run.font.size = Pt(11)
    run.font.name = "Calibri"
    run.font.color.rgb = RGBColor(0x26, 0x4F, 0x78)


def _add_entry_heading(doc: Document, parts: list[str]) -> None:
    cleaned = [part.strip() if part and part.strip() else "[unknown]" for part in parts]
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(1)
    run = para.add_run(" | ".join(cleaned))
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.name = "Calibri"


def _add_bullets(doc: Document, bullets: list[str]) -> None:
    for bullet in bullets:
        if not bullet.strip():
            continue
        para = doc.add_paragraph(style="List Bullet")
        para.paragraph_format.space_after = Pt(0)
        _set_run(para.add_run(bullet.strip()), size=10.5)


def _add_text(doc: Document, text: str) -> None:
    if not text.strip():
        return
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(2)
    _set_run(para.add_run(text.strip()), size=10.5)


def _set_run(run, size: float) -> None:
    run.font.size = Pt(size)
    run.font.name = "Calibri"


def _warn_missing(kind: str, values: dict, fields: list[str]) -> None:
    missing = [field for field in fields if not values.get(field)]
    if missing:
        logger.warning("%s entry missing fields: %s", kind, ", ".join(missing))
